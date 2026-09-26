from __future__ import annotations

import csv
import json
import math
import re
import unicodedata
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

RAW_FILE = (
    PROJECT_ROOT
    / "01_raw_data"
    / "MaxT_MinT_1981_2024_raw.xlsx"
)

STATION_MASTER_FILE = (
    PROJECT_ROOT
    / "02_metadata"
    / "station_master_v2.csv"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_intermediate"
    / "old_xlsx_daily_standardized.parquet"
)

TEMP_OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_intermediate"
    / ".old_xlsx_daily_standardized.tmp.parquet"
)

REPORT_DIRECTORY = PROJECT_ROOT / "05_qc_reports"
ADMIN_DIRECTORY = PROJECT_ROOT / "00_admin"

SHEET_NAME = "MaxT & MinT"
HEADER_ROW = 2
FIRST_DATA_ROW = 3
BATCH_SIZE = 50_000

EXPECTED_RECORDS = 548_213
EXPECTED_STATIONS = 42

MISSING_CODES = {
    "",
    "****",
    "**",
    "na",
    "n/a",
    "nan",
    "none",
    "null",
}


SCHEMA = pa.schema(
    [
        pa.field("source_dataset", pa.string()),
        pa.field("source_file", pa.string()),
        pa.field("source_sheet", pa.string()),
        pa.field("source_row", pa.int32()),

        pa.field("station_name_raw", pa.string()),
        pa.field("station_name_key", pa.string()),
        pa.field("station_uid", pa.string()),
        pa.field("station_id", pa.string()),
        pa.field("station_name_official", pa.string()),
        pa.field("latitude", pa.float64()),
        pa.field("longitude", pa.float64()),
        pa.field("station_match_status", pa.string()),

        pa.field("year_raw", pa.string()),
        pa.field("month_raw", pa.string()),
        pa.field("day_raw", pa.string()),
        pa.field("year", pa.int16()),
        pa.field("month", pa.int8()),
        pa.field("day", pa.int8()),
        pa.field("date", pa.date32()),
        pa.field("date_parse_status", pa.string()),

        pa.field("tmax_raw", pa.string()),
        pa.field("tmax_old_xlsx", pa.float32()),
        pa.field("tmax_parse_status", pa.string()),

        pa.field("tmin_raw", pa.string()),
        pa.field("tmin_old_xlsx", pa.float32()),
        pa.field("tmin_parse_status", pa.string()),

        pa.field("structural_status", pa.string()),
    ]
)


def normalize_name(value: Any) -> str:
    """Create a punctuation- and case-insensitive station key."""
    if value is None:
        return ""

    text = unicodedata.normalize(
        "NFKD",
        str(value).strip(),
    )

    text = (
        text.encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )

    return re.sub(r"[^a-z0-9]+", "", text)


def raw_text(value: Any) -> str:
    """Preserve an Excel value as compact text."""
    if value is None:
        return ""

    if isinstance(value, float):
        if math.isnan(value):
            return ""

        if value.is_integer():
            return str(int(value))

    return str(value).strip()


def clean_station_id(value: Any) -> str:
    """Remove Excel-style decimal suffixes from station IDs."""
    text = raw_text(value)

    if text.endswith(".0"):
        text = text[:-2]

    return text


def parse_integer(
    value: Any,
) -> tuple[int | None, str]:
    """Parse a date component without correcting it."""
    text = raw_text(value)

    if text.lower() in MISSING_CODES:
        return None, "missing"

    try:
        number = float(text)
    except (TypeError, ValueError):
        return None, "non_numeric"

    if not math.isfinite(number):
        return None, "non_numeric"

    if not number.is_integer():
        return None, "non_integer"

    return int(number), "valid_integer"


def parse_date(
    year_value: Any,
    month_value: Any,
    day_value: Any,
) -> tuple[
    int | None,
    int | None,
    int | None,
    date | None,
    str,
]:
    """Build a valid date or return an explicit failure status."""
    year, year_status = parse_integer(year_value)
    month, month_status = parse_integer(month_value)
    day, day_status = parse_integer(day_value)

    statuses = {
        year_status,
        month_status,
        day_status,
    }

    if "missing" in statuses:
        return (
            year,
            month,
            day,
            None,
            "missing_date_component",
        )

    if statuses != {"valid_integer"}:
        return (
            year,
            month,
            day,
            None,
            "invalid_date_component",
        )

    try:
        parsed_date = date(
            int(year),
            int(month),
            int(day),
        )
    except ValueError:
        return (
            year,
            month,
            day,
            None,
            "invalid_calendar_date",
        )

    return (
        year,
        month,
        day,
        parsed_date,
        "valid",
    )


def parse_temperature(
    value: Any,
) -> tuple[float | None, str]:
    """Parse temperature while preserving missing/error status."""
    text = raw_text(value)

    if value is None or text == "":
        return None, "missing_cell"

    if text.lower() in MISSING_CODES:
        return None, "missing_code"

    try:
        number = float(text)
    except (TypeError, ValueError):
        return None, "non_numeric"

    if not math.isfinite(number):
        return None, "non_numeric"

    return number, "numeric"


def load_old_station_mapping() -> dict[str, dict[str, Any]]:
    """Load station identities for stations represented in the old XLSX."""
    with STATION_MASTER_FILE.open(
        "r",
        encoding="utf-8",
        newline="",
    ) as file_handle:
        rows = list(csv.DictReader(file_handle))

    mapping: dict[str, dict[str, Any]] = {}

    for row in rows:
        old_name = row[
            "old_xlsx_station_name"
        ].strip()

        if not old_name:
            continue

        key = normalize_name(old_name)

        if key in mapping:
            raise ValueError(
                "Duplicate old-XLSX station key: "
                f"{old_name!r}"
            )

        if (
            row["metadata_status"].strip()
            != "official_metadata_available"
        ):
            raise ValueError(
                "An old-XLSX station lacks official "
                f"metadata: {old_name!r}"
            )

        mapping[key] = {
            "station_uid":
                row["station_uid"].strip(),
            "station_id":
                clean_station_id(
                    row["station_id"]
                ),
            "station_name_official":
                row[
                    "official_station_name"
                ].strip(),
            "latitude":
                float(row["latitude"]),
            "longitude":
                float(row["longitude"]),
        }

    if len(mapping) != EXPECTED_STATIONS:
        raise ValueError(
            f"Expected {EXPECTED_STATIONS} old-XLSX "
            f"station mappings; found {len(mapping)}."
        )

    return mapping


def empty_batch() -> dict[str, list[Any]]:
    return {
        field.name: []
        for field in SCHEMA
    }


def append_record(
    batch: dict[str, list[Any]],
    record: dict[str, Any],
) -> None:
    for field_name in batch:
        batch[field_name].append(
            record[field_name]
        )


def write_batch(
    writer: pq.ParquetWriter,
    batch: dict[str, list[Any]],
) -> int:
    row_count = len(batch["source_row"])

    if row_count == 0:
        return 0

    table = pa.Table.from_pydict(
        batch,
        schema=SCHEMA,
    )

    writer.write_table(table)

    return row_count


def write_issue_csv(
    path: Path,
    records: list[dict[str, Any]],
    columns: list[str],
) -> None:
    with path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as file_handle:
        writer = csv.DictWriter(
            file_handle,
            fieldnames=columns,
        )

        writer.writeheader()
        writer.writerows(records)


def main() -> None:
    for required_path in [
        RAW_FILE,
        STATION_MASTER_FILE,
    ]:
        if not required_path.exists():
            raise FileNotFoundError(
                f"Required file not found: {required_path}"
            )

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    REPORT_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True,
    )

    ADMIN_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True,
    )

    if TEMP_OUTPUT_FILE.exists():
        TEMP_OUTPUT_FILE.unlink()

    station_mapping = load_old_station_mapping()

    workbook = load_workbook(
        RAW_FILE,
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    writer: pq.ParquetWriter | None = None

    date_status_counts: Counter[str] = Counter()
    tmax_status_counts: Counter[str] = Counter()
    tmin_status_counts: Counter[str] = Counter()
    structural_status_counts: Counter[str] = Counter()
    station_record_counts: Counter[str] = Counter()

    unique_station_names: set[str] = set()
    unique_station_uids: set[str] = set()

    invalid_date_records: list[dict[str, Any]] = []
    unmatched_station_records: list[dict[str, Any]] = []
    unparsed_temperature_records: list[
        dict[str, Any]
    ] = []

    source_rows_examined = 0
    blank_rows_skipped = 0
    records_written = 0
    unmatched_row_count = 0

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet {SHEET_NAME!r} not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[SHEET_NAME]

        observed_headers = [
            worksheet.cell(
                row=HEADER_ROW,
                column=column,
            ).value
            for column in range(1, 7)
        ]

        expected_headers = [
            "Station",
            "YEAR",
            "Month",
            "DATE",
            "Max.T",
            "Min.T",
        ]

        observed_keys = [
            normalize_name(value)
            for value in observed_headers
        ]

        expected_keys = [
            normalize_name(value)
            for value in expected_headers
        ]

        if observed_keys != expected_keys:
            raise ValueError(
                "Unexpected workbook headers.\n"
                f"Expected: {expected_headers}\n"
                f"Observed: {observed_headers}"
            )

        writer = pq.ParquetWriter(
            TEMP_OUTPUT_FILE,
            SCHEMA,
            compression="zstd",
            use_dictionary=True,
        )

        batch = empty_batch()

        for source_row, values in enumerate(
            worksheet.iter_rows(
                min_row=FIRST_DATA_ROW,
                min_col=1,
                max_col=6,
                values_only=True,
            ),
            start=FIRST_DATA_ROW,
        ):
            source_rows_examined += 1

            if all(
                value is None
                or (
                    isinstance(value, str)
                    and value.strip() == ""
                )
                for value in values
            ):
                blank_rows_skipped += 1
                continue

            (
                station_value,
                year_value,
                month_value,
                day_value,
                tmax_value,
                tmin_value,
            ) = values

            station_name_raw = raw_text(
                station_value
            )

            station_name_key = normalize_name(
                station_name_raw
            )

            unique_station_names.add(
                station_name_raw
            )

            station = station_mapping.get(
                station_name_key
            )

            if station is None:
                station_uid = ""
                station_id = ""
                station_name_official = ""
                latitude = None
                longitude = None
                station_match_status = "unmatched"

                unmatched_row_count += 1

                unmatched_station_records.append(
                    {
                        "source_row": source_row,
                        "station_name_raw":
                            station_name_raw,
                        "station_name_key":
                            station_name_key,
                    }
                )

            else:
                station_uid = station[
                    "station_uid"
                ]

                station_id = station[
                    "station_id"
                ]

                station_name_official = station[
                    "station_name_official"
                ]

                latitude = station["latitude"]
                longitude = station["longitude"]
                station_match_status = "matched"

                unique_station_uids.add(
                    station_uid
                )

                station_record_counts[
                    station_uid
                ] += 1

            (
                year,
                month,
                day,
                parsed_date,
                date_parse_status,
            ) = parse_date(
                year_value,
                month_value,
                day_value,
            )

            (
                tmax_numeric,
                tmax_parse_status,
            ) = parse_temperature(
                tmax_value
            )

            (
                tmin_numeric,
                tmin_parse_status,
            ) = parse_temperature(
                tmin_value
            )

            date_status_counts[
                date_parse_status
            ] += 1

            tmax_status_counts[
                tmax_parse_status
            ] += 1

            tmin_status_counts[
                tmin_parse_status
            ] += 1

            has_unparsed_temperature = (
                tmax_parse_status == "non_numeric"
                or tmin_parse_status == "non_numeric"
            )

            if (
                station_match_status == "matched"
                and date_parse_status == "valid"
                and not has_unparsed_temperature
            ):
                structural_status = "ready_for_qc"
            else:
                structural_status = (
                    "review_required"
                )

            structural_status_counts[
                structural_status
            ] += 1

            common_issue_fields = {
                "source_row": source_row,
                "station_name_raw":
                    station_name_raw,
                "station_uid": station_uid,
                "year_raw": raw_text(
                    year_value
                ),
                "month_raw": raw_text(
                    month_value
                ),
                "day_raw": raw_text(
                    day_value
                ),
                "tmax_raw": raw_text(
                    tmax_value
                ),
                "tmin_raw": raw_text(
                    tmin_value
                ),
            }

            if date_parse_status != "valid":
                invalid_date_records.append(
                    {
                        **common_issue_fields,
                        "date_parse_status":
                            date_parse_status,
                    }
                )

            if has_unparsed_temperature:
                unparsed_temperature_records.append(
                    {
                        **common_issue_fields,
                        "tmax_parse_status":
                            tmax_parse_status,
                        "tmin_parse_status":
                            tmin_parse_status,
                    }
                )

            record = {
                "source_dataset":
                    "old_xlsx_1981_2024",
                "source_file": RAW_FILE.name,
                "source_sheet": SHEET_NAME,
                "source_row": source_row,

                "station_name_raw":
                    station_name_raw,
                "station_name_key":
                    station_name_key,
                "station_uid": station_uid,
                "station_id": station_id,
                "station_name_official":
                    station_name_official,
                "latitude": latitude,
                "longitude": longitude,
                "station_match_status":
                    station_match_status,

                "year_raw": raw_text(
                    year_value
                ),
                "month_raw": raw_text(
                    month_value
                ),
                "day_raw": raw_text(
                    day_value
                ),
                "year": year,
                "month": month,
                "day": day,
                "date": parsed_date,
                "date_parse_status":
                    date_parse_status,

                "tmax_raw": raw_text(
                    tmax_value
                ),
                "tmax_old_xlsx":
                    tmax_numeric,
                "tmax_parse_status":
                    tmax_parse_status,

                "tmin_raw": raw_text(
                    tmin_value
                ),
                "tmin_old_xlsx":
                    tmin_numeric,
                "tmin_parse_status":
                    tmin_parse_status,

                "structural_status":
                    structural_status,
            }

            append_record(batch, record)

            if (
                len(batch["source_row"])
                >= BATCH_SIZE
            ):
                records_written += write_batch(
                    writer,
                    batch,
                )

                print(
                    f"Written {records_written:,} records..."
                )

                batch = empty_batch()

        if batch["source_row"]:
            records_written += write_batch(
                writer,
                batch,
            )

    except Exception:
        if writer is not None:
            writer.close()
            writer = None

        if TEMP_OUTPUT_FILE.exists():
            TEMP_OUTPUT_FILE.unlink()

        raise

    finally:
        if writer is not None:
            writer.close()

        workbook.close()

    if OUTPUT_FILE.exists():
        OUTPUT_FILE.unlink()

    TEMP_OUTPUT_FILE.replace(OUTPUT_FILE)

    invalid_date_columns = [
        "source_row",
        "station_name_raw",
        "station_uid",
        "year_raw",
        "month_raw",
        "day_raw",
        "tmax_raw",
        "tmin_raw",
        "date_parse_status",
    ]

    unmatched_columns = [
        "source_row",
        "station_name_raw",
        "station_name_key",
    ]

    unparsed_columns = [
        "source_row",
        "station_name_raw",
        "station_uid",
        "year_raw",
        "month_raw",
        "day_raw",
        "tmax_raw",
        "tmin_raw",
        "tmax_parse_status",
        "tmin_parse_status",
    ]

    write_issue_csv(
        REPORT_DIRECTORY
        / "step3a_invalid_date_rows.csv",
        invalid_date_records,
        invalid_date_columns,
    )

    write_issue_csv(
        REPORT_DIRECTORY
        / "step3a_unmatched_station_rows.csv",
        unmatched_station_records,
        unmatched_columns,
    )

    write_issue_csv(
        REPORT_DIRECTORY
        / "step3a_unparsed_temperature_rows.csv",
        unparsed_temperature_records,
        unparsed_columns,
    )

    station_count_rows = [
        {
            "station_uid": station_uid,
            "record_count": count,
        }
        for station_uid, count
        in sorted(station_record_counts.items())
    ]

    write_issue_csv(
        REPORT_DIRECTORY
        / "step3a_station_record_counts.csv",
        station_count_rows,
        [
            "station_uid",
            "record_count",
        ],
    )

    parquet_file = pq.ParquetFile(
        OUTPUT_FILE
    )

    parquet_rows = (
        parquet_file.metadata.num_rows
    )

    summary = {
        "created_utc":
            datetime.now(
                timezone.utc
            ).isoformat(),
        "source_file": str(RAW_FILE),
        "source_sheet": SHEET_NAME,
        "source_rows_examined":
            source_rows_examined,
        "blank_rows_skipped":
            blank_rows_skipped,
        "records_written":
            records_written,
        "parquet_rows":
            parquet_rows,
        "unique_raw_station_names":
            len(unique_station_names),
        "unique_matched_station_uids":
            len(unique_station_uids),
        "unmatched_station_rows":
            unmatched_row_count,
        "invalid_date_rows":
            len(invalid_date_records),
        "unparsed_temperature_rows":
            len(
                unparsed_temperature_records
            ),
        "date_parse_status_counts":
            dict(date_status_counts),
        "tmax_parse_status_counts":
            dict(tmax_status_counts),
        "tmin_parse_status_counts":
            dict(tmin_status_counts),
        "structural_status_counts":
            dict(structural_status_counts),
        "output_file": str(OUTPUT_FILE),
        "output_size_bytes":
            OUTPUT_FILE.stat().st_size,
    }

    summary_path = (
        REPORT_DIRECTORY
        / "step3a_ingestion_summary.json"
    )

    summary_path.write_text(
        json.dumps(
            summary,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    schema_path = (
        ADMIN_DIRECTORY
        / "step3a_parquet_schema.json"
    )

    schema_path.write_text(
        json.dumps(
            {
                field.name: str(field.type)
                for field in SCHEMA
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    report_lines = [
        "STEP 3A: OLD XLSX INGESTION SUMMARY",
        "=" * 40,
        (
            "Source rows examined: "
            f"{source_rows_examined:,}"
        ),
        (
            "Blank rows skipped: "
            f"{blank_rows_skipped:,}"
        ),
        (
            "Records written: "
            f"{records_written:,}"
        ),
        (
            "Parquet rows: "
            f"{parquet_rows:,}"
        ),
        (
            "Unique raw station names: "
            f"{len(unique_station_names)}"
        ),
        (
            "Unique matched station UIDs: "
            f"{len(unique_station_uids)}"
        ),
        (
            "Unmatched station rows: "
            f"{unmatched_row_count:,}"
        ),
        (
            "Invalid date rows: "
            f"{len(invalid_date_records):,}"
        ),
        (
            "Unparsed temperature rows: "
            f"{len(unparsed_temperature_records):,}"
        ),
        "",
        "Date parse statuses:",
        *[
            f"- {key}: {value:,}"
            for key, value
            in sorted(
                date_status_counts.items()
            )
        ],
        "",
        "Tmax parse statuses:",
        *[
            f"- {key}: {value:,}"
            for key, value
            in sorted(
                tmax_status_counts.items()
            )
        ],
        "",
        "Tmin parse statuses:",
        *[
            f"- {key}: {value:,}"
            for key, value
            in sorted(
                tmin_status_counts.items()
            )
        ],
        "",
        "Structural statuses:",
        *[
            f"- {key}: {value:,}"
            for key, value
            in sorted(
                structural_status_counts.items()
            )
        ],
        "",
        f"Output: {OUTPUT_FILE}",
    ]

    report_path = (
        REPORT_DIRECTORY
        / "step3a_ingestion_report.txt"
    )

    report_path.write_text(
        "\n".join(report_lines) + "\n",
        encoding="utf-8",
    )

    print()
    print("\n".join(report_lines))

    failures: list[str] = []

    if records_written != EXPECTED_RECORDS:
        failures.append(
            f"Expected {EXPECTED_RECORDS:,} records "
            f"but wrote {records_written:,}."
        )

    if parquet_rows != records_written:
        failures.append(
            "Parquet row count differs from "
            "records written."
        )

    if (
        len(unique_station_names)
        != EXPECTED_STATIONS
    ):
        failures.append(
            f"Expected {EXPECTED_STATIONS} raw "
            "station names but found "
            f"{len(unique_station_names)}."
        )

    if (
        len(unique_station_uids)
        != EXPECTED_STATIONS
    ):
        failures.append(
            f"Expected {EXPECTED_STATIONS} matched "
            "station UIDs but found "
            f"{len(unique_station_uids)}."
        )

    if unmatched_row_count != 0:
        failures.append(
            "One or more rows have unmatched stations."
        )

    if failures:
        raise SystemExit(
            "\nSTEP 3A FAILED:\n- "
            + "\n- ".join(failures)
        )

    print("\nSTEP 3A PASSED.")


if __name__ == "__main__":
    main()

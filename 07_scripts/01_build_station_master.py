from __future__ import annotations

import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from rapidfuzz import fuzz, process


PROJECT_ROOT = Path(__file__).resolve().parents[1]

TEMP_FILE = (
    PROJECT_ROOT
    / "01_raw_data"
    / "MaxT_MinT_1981_2024_raw.xlsx"
)

META_FILE = (
    PROJECT_ROOT
    / "01_raw_data"
    / "All_station_BMD_raw.xlsx"
)

ALIAS_FILE = (
    PROJECT_ROOT
    / "02_metadata"
    / "station_aliases.csv"
)

TEMP_SHEET = "MaxT & MinT"
META_SHEET = "Station_Info"

TEMP_HEADER_ROW = 2
TEMP_STATION_COLUMN = 1
TEMP_FIRST_DATA_ROW = 3

EXPECTED_TEMP_STATIONS = 42
EXPECTED_METADATA_STATIONS = 46
EXPECTED_METADATA_ONLY = 4


def normalize_name(value: object) -> str:
    """
    Create a comparison key from a station name.

    Capitalization, spaces, periods, apostrophes, underscores
    and other punctuation are removed.
    """
    text = unicodedata.normalize(
        "NFKD",
        str(value).strip(),
    )

    text = (
        text.encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )

    return re.sub(r"[^a-z0-9]+", "", text)


def clean_station_id(value: object) -> str:
    """Convert Excel station IDs into clean strings."""
    if pd.isna(value):
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def read_temperature_station_names() -> list[str]:
    """
    Stream only the station column from the large workbook.

    This avoids loading all 548,000 rows into memory.
    """
    workbook = load_workbook(
        TEMP_FILE,
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    try:
        if TEMP_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Sheet {TEMP_SHEET!r} was not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[TEMP_SHEET]

        header = worksheet.cell(
            row=TEMP_HEADER_ROW,
            column=TEMP_STATION_COLUMN,
        ).value

        if normalize_name(header) != "station":
            raise ValueError(
                f"Expected 'Station' at row "
                f"{TEMP_HEADER_ROW}, column "
                f"{TEMP_STATION_COLUMN}; found {header!r}."
            )

        station_names: set[str] = set()

        for (value,) in worksheet.iter_rows(
            min_row=TEMP_FIRST_DATA_ROW,
            min_col=TEMP_STATION_COLUMN,
            max_col=TEMP_STATION_COLUMN,
            values_only=True,
        ):
            if value is None:
                continue

            name = str(value).strip()

            if name:
                station_names.add(name)

        return sorted(
            station_names,
            key=str.casefold,
        )

    finally:
        workbook.close()


def read_metadata() -> pd.DataFrame:
    """Read and standardize the BMD metadata workbook."""
    metadata = pd.read_excel(
        META_FILE,
        sheet_name=META_SHEET,
        engine="openpyxl",
    )

    required_columns = {
        "St_name",
        "St_ID",
        "Lat",
        "Lon",
    }

    missing_columns = (
        required_columns.difference(metadata.columns)
    )

    if missing_columns:
        raise ValueError(
            "Missing metadata columns: "
            f"{sorted(missing_columns)}"
        )

    metadata = metadata.loc[
        :,
        ["St_name", "St_ID", "Lat", "Lon"],
    ].copy()

    metadata.columns = [
        "official_station_name",
        "station_id",
        "latitude",
        "longitude",
    ]

    metadata = metadata.dropna(how="all")

    metadata["official_station_name"] = (
        metadata["official_station_name"]
        .astype(str)
        .str.strip()
    )

    metadata["station_id"] = (
        metadata["station_id"]
        .map(clean_station_id)
    )

    metadata["latitude"] = pd.to_numeric(
        metadata["latitude"],
        errors="coerce",
    )

    metadata["longitude"] = pd.to_numeric(
        metadata["longitude"],
        errors="coerce",
    )

    metadata["official_name_key"] = (
        metadata["official_station_name"]
        .map(normalize_name)
    )

    return metadata


def validate_metadata(
    metadata: pd.DataFrame,
) -> None:
    """Stop the workflow if metadata problems are found."""
    problems: list[str] = []

    if metadata["station_id"].eq("").any():
        problems.append(
            "At least one station ID is blank."
        )

    if metadata["station_id"].duplicated().any():
        duplicate_ids = metadata.loc[
            metadata["station_id"].duplicated(
                keep=False
            ),
            "station_id",
        ].tolist()

        problems.append(
            f"Duplicate station IDs: {duplicate_ids}"
        )

    if metadata["official_name_key"].duplicated().any():
        duplicate_names = metadata.loc[
            metadata["official_name_key"].duplicated(
                keep=False
            ),
            "official_station_name",
        ].tolist()

        problems.append(
            "Duplicate normalized station names: "
            f"{duplicate_names}"
        )

    if metadata[
        ["latitude", "longitude"]
    ].isna().any().any():
        problems.append(
            "Missing or nonnumeric coordinates found."
        )

    invalid_coordinates = metadata.loc[
        ~metadata["latitude"].between(20.0, 27.5)
        | ~metadata["longitude"].between(88.0, 93.5)
    ]

    if not invalid_coordinates.empty:
        station_list = ", ".join(
            invalid_coordinates[
                "official_station_name"
            ]
        )

        problems.append(
            "Coordinates outside broad Bangladesh "
            f"bounds: {station_list}"
        )

    if problems:
        raise ValueError(
            "Metadata validation failed:\n- "
            + "\n- ".join(problems)
        )


def read_aliases(
    metadata: pd.DataFrame,
) -> dict[str, str]:
    """Read the manually verified station aliases."""
    aliases = pd.read_csv(ALIAS_FILE)

    required_columns = {
        "raw_station_name",
        "official_station_name",
        "reason",
    }

    missing_columns = (
        required_columns.difference(aliases.columns)
    )

    if missing_columns:
        raise ValueError(
            "Missing alias-file columns: "
            f"{sorted(missing_columns)}"
        )

    official_keys = set(
        metadata["official_name_key"]
    )

    alias_map: dict[str, str] = {}

    for row in aliases.itertuples(index=False):
        raw_key = normalize_name(
            row.raw_station_name
        )

        official_key = normalize_name(
            row.official_station_name
        )

        if official_key not in official_keys:
            raise ValueError(
                f"Alias target "
                f"{row.official_station_name!r} "
                "does not exist in the BMD metadata."
            )

        if (
            raw_key in alias_map
            and alias_map[raw_key] != official_key
        ):
            raise ValueError(
                "Conflicting alias entries for "
                f"{row.raw_station_name!r}."
            )

        alias_map[raw_key] = official_key

    return alias_map


def build_crosswalk(
    raw_names: list[str],
    metadata: pd.DataFrame,
    aliases: dict[str, str],
) -> pd.DataFrame:
    """Match each raw temperature name to BMD metadata."""
    metadata_by_key = (
        metadata
        .set_index("official_name_key")
        .to_dict("index")
    )

    official_keys = list(metadata_by_key)

    records: list[dict] = []

    for raw_name in raw_names:
        raw_key = normalize_name(raw_name)

        matched_key: str | None = None
        method = "unmatched"
        match_score: float | None = None

        if raw_key in aliases:
            matched_key = aliases[raw_key]
            method = "manual_alias"
            match_score = 100.0

        elif raw_key in metadata_by_key:
            matched_key = raw_key
            method = "normalized_exact"
            match_score = 100.0

        suggestion = ""
        suggestion_score: float | None = None

        if matched_key is None:
            fuzzy_result = process.extractOne(
                raw_key,
                official_keys,
                scorer=fuzz.ratio,
            )

            if fuzzy_result is not None:
                (
                    suggested_key,
                    suggestion_score,
                    _,
                ) = fuzzy_result

                suggestion = metadata_by_key[
                    suggested_key
                ]["official_station_name"]

        if matched_key is not None:
            station = metadata_by_key[matched_key]

            records.append(
                {
                    "raw_station_name": raw_name,
                    "raw_name_key": raw_key,
                    "official_station_name": station[
                        "official_station_name"
                    ],
                    "station_id": station[
                        "station_id"
                    ],
                    "latitude": station[
                        "latitude"
                    ],
                    "longitude": station[
                        "longitude"
                    ],
                    "match_method": method,
                    "match_score": match_score,
                    "suggested_official_name": "",
                    "suggestion_score": "",
                    "action_required": False,
                }
            )

        else:
            records.append(
                {
                    "raw_station_name": raw_name,
                    "raw_name_key": raw_key,
                    "official_station_name": "",
                    "station_id": "",
                    "latitude": pd.NA,
                    "longitude": pd.NA,
                    "match_method": "unmatched",
                    "match_score": pd.NA,
                    "suggested_official_name": suggestion,
                    "suggestion_score": suggestion_score,
                    "action_required": True,
                }
            )

    return pd.DataFrame(records)


def main() -> None:
    required_files = [
        TEMP_FILE,
        META_FILE,
        ALIAS_FILE,
    ]

    for path in required_files:
        if not path.exists():
            raise FileNotFoundError(
                f"Required file not found: {path}"
            )

    raw_names = read_temperature_station_names()

    metadata = read_metadata()
    validate_metadata(metadata)

    aliases = read_aliases(metadata)

    crosswalk = build_crosswalk(
        raw_names,
        metadata,
        aliases,
    )

    unmatched = crosswalk.loc[
        crosswalk["action_required"]
    ].copy()

    matched = crosswalk.loc[
        ~crosswalk["action_required"]
    ].copy()

    duplicate_matches = matched.loc[
        matched["station_id"].duplicated(
            keep=False
        )
    ].copy()

    raw_names_by_id = (
        matched
        .groupby("station_id")[
            "raw_station_name"
        ]
        .agg(
            lambda values: "; ".join(
                sorted(
                    values,
                    key=str.casefold,
                )
            )
        )
        .to_dict()
    )

    methods_by_id = (
        matched
        .groupby("station_id")[
            "match_method"
        ]
        .agg(
            lambda values: "; ".join(
                sorted(set(values))
            )
        )
        .to_dict()
    )

    station_master = metadata.copy()

    station_master[
        "has_temperature_data"
    ] = station_master["station_id"].isin(
        matched["station_id"]
    )

    station_master[
        "raw_temperature_name"
    ] = (
        station_master["station_id"]
        .map(raw_names_by_id)
        .fillna("")
    )

    station_master[
        "match_method"
    ] = (
        station_master["station_id"]
        .map(methods_by_id)
        .fillna("")
    )

    station_master["elevation_m"] = pd.NA
    station_master["metadata_notes"] = ""

    station_master = station_master[
        [
            "station_id",
            "official_station_name",
            "raw_temperature_name",
            "latitude",
            "longitude",
            "elevation_m",
            "has_temperature_data",
            "match_method",
            "metadata_notes",
        ]
    ].sort_values("station_id")

    temperature_station_master = (
        station_master.loc[
            station_master[
                "has_temperature_data"
            ]
        ].copy()
    )

    metadata_only = station_master.loc[
        ~station_master[
            "has_temperature_data"
        ]
    ].copy()

    output_directory = (
        PROJECT_ROOT / "02_metadata"
    )

    report_directory = (
        PROJECT_ROOT / "05_qc_reports"
    )

    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    report_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    crosswalk.to_csv(
        output_directory
        / "station_name_crosswalk.csv",
        index=False,
    )

    station_master.to_csv(
        output_directory
        / "station_master.csv",
        index=False,
    )

    temperature_station_master.to_csv(
        output_directory
        / "temperature_station_master.csv",
        index=False,
    )

    metadata_only.to_csv(
        output_directory
        / "metadata_only_stations.csv",
        index=False,
    )

    unmatched.to_csv(
        output_directory
        / "unmatched_temperature_stations.csv",
        index=False,
    )

    duplicate_matches.to_csv(
        output_directory
        / "duplicate_station_matches.csv",
        index=False,
    )

    with pd.ExcelWriter(
        output_directory
        / "station_identity_tables.xlsx",
        engine="openpyxl",
    ) as writer:
        station_master.to_excel(
            writer,
            sheet_name="station_master",
            index=False,
        )

        crosswalk.to_excel(
            writer,
            sheet_name="crosswalk",
            index=False,
        )

        metadata_only.to_excel(
            writer,
            sheet_name="metadata_only",
            index=False,
        )

        unmatched.to_excel(
            writer,
            sheet_name="unmatched",
            index=False,
        )

    summary = {
        "created_utc": datetime.now(
            timezone.utc
        ).isoformat(),
        "temperature_unique_station_names":
            len(raw_names),
        "metadata_station_records":
            len(metadata),
        "matched_temperature_stations":
            len(matched),
        "unmatched_temperature_stations":
            len(unmatched),
        "metadata_only_stations":
            len(metadata_only),
        "duplicate_station_id_matches":
            len(duplicate_matches),
        "metadata_only_station_names":
            metadata_only[
                "official_station_name"
            ].tolist(),
    }

    summary_path = (
        report_directory
        / "step2_station_matching_summary.json"
    )

    summary_path.write_text(
        json.dumps(
            summary,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    report_lines = [
        "STEP 2: STATION IDENTITY SUMMARY",
        "=" * 36,
        (
            "Temperature station names: "
            f"{len(raw_names)}"
        ),
        (
            "Metadata station records: "
            f"{len(metadata)}"
        ),
        (
            "Matched temperature stations: "
            f"{len(matched)}"
        ),
        (
            "Unmatched temperature stations: "
            f"{len(unmatched)}"
        ),
        (
            "Metadata-only stations: "
            f"{len(metadata_only)}"
        ),
        (
            "Duplicate station-ID matches: "
            f"{len(duplicate_matches)}"
        ),
        "",
        "Metadata-only station names:",
        *[
            f"- {name}"
            for name in metadata_only[
                "official_station_name"
            ]
        ],
    ]

    report_path = (
        report_directory
        / "step2_station_matching_report.txt"
    )

    report_path.write_text(
        "\n".join(report_lines) + "\n",
        encoding="utf-8",
    )

    print("\n".join(report_lines))

    failures: list[str] = []

    if len(raw_names) != EXPECTED_TEMP_STATIONS:
        failures.append(
            f"Expected {EXPECTED_TEMP_STATIONS} "
            "temperature stations, but found "
            f"{len(raw_names)}."
        )

    if len(metadata) != EXPECTED_METADATA_STATIONS:
        failures.append(
            f"Expected {EXPECTED_METADATA_STATIONS} "
            "metadata stations, but found "
            f"{len(metadata)}."
        )

    if not unmatched.empty:
        failures.append(
            "Some temperature station names "
            "remain unmatched."
        )

    if not duplicate_matches.empty:
        failures.append(
            "More than one raw station name "
            "maps to the same station ID."
        )

    if len(metadata_only) != EXPECTED_METADATA_ONLY:
        failures.append(
            f"Expected {EXPECTED_METADATA_ONLY} "
            "metadata-only stations, but found "
            f"{len(metadata_only)}."
        )

    if failures:
        raise SystemExit(
            "\nSTEP 2 FAILED:\n- "
            + "\n- ".join(failures)
        )

    print("\nSTEP 2 PASSED.")


if __name__ == "__main__":
    main()

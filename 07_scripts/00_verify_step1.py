from __future__ import annotations

import hashlib
import json
import platform
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILES = {
    "temperature": (
        PROJECT_ROOT
        / "01_raw_data"
        / "MaxT_MinT_1981_2024_raw.xlsx"
    ),
    "station_metadata": (
        PROJECT_ROOT
        / "01_raw_data"
        / "All_station_BMD_raw.xlsx"
    ),
}


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    """Calculate SHA-256 without loading the complete file into memory."""
    digest = hashlib.sha256()

    with path.open("rb") as file_handle:
        while chunk := file_handle.read(chunk_size):
            digest.update(chunk)

    return digest.hexdigest()


def inspect_workbook(path: Path) -> dict:
    """Read workbook structure without loading all cell data."""
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )

    try:
        sheets = []

        for worksheet in workbook.worksheets:
            preview = []

            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(3, worksheet.max_row),
                values_only=True,
            ):
                preview.append(list(row[:10]))

            sheets.append(
                {
                    "sheet_name": worksheet.title,
                    "reported_rows": worksheet.max_row,
                    "reported_columns": worksheet.max_column,
                    "first_rows_preview": preview,
                }
            )

        return {"sheets": sheets}

    finally:
        workbook.close()


def main() -> None:
    report = {
        "verification_time_utc": datetime.now(
            timezone.utc
        ).isoformat(),
        "python_version": sys.version,
        "platform": platform.platform(),
        "files": {},
    }

    all_checks_passed = True

    for label, path in FILES.items():
        print(f"\nChecking {label}: {path}")

        if not path.exists():
            print("  ERROR: File does not exist.")
            report["files"][label] = {
                "path": str(path),
                "exists": False,
            }
            all_checks_passed = False
            continue

        if path.stat().st_size == 0:
            print("  ERROR: File is empty.")
            all_checks_passed = False

        try:
            workbook_information = inspect_workbook(path)
            workbook_readable = True
        except Exception as exc:
            workbook_information = {"error": repr(exc)}
            workbook_readable = False
            all_checks_passed = False

        file_report = {
            "path": str(path),
            "exists": True,
            "size_bytes": path.stat().st_size,
            "sha256": sha256_file(path),
            "workbook_readable": workbook_readable,
            **workbook_information,
        }

        report["files"][label] = file_report

        print(f"  Size: {file_report['size_bytes']:,} bytes")
        print(f"  SHA-256: {file_report['sha256']}")
        print(f"  Workbook readable: {workbook_readable}")

        if workbook_readable:
            for sheet in file_report["sheets"]:
                print(
                    "  Sheet:",
                    sheet["sheet_name"],
                    "| rows:",
                    sheet["reported_rows"],
                    "| columns:",
                    sheet["reported_columns"],
                )

    output_path = (
        PROJECT_ROOT
        / "00_admin"
        / "step1_verification_report.json"
    )

    output_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print(f"\nVerification report written to: {output_path}")

    if not all_checks_passed:
        raise SystemExit(
            "\nSTEP 1 FAILED. Correct the errors before continuing."
        )

    print("\nSTEP 1 PASSED.")


if __name__ == "__main__":
    main()
